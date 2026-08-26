"""
main.py — Serviço FastAPI de Exportação DOCX de Alta Fidelidade
"The Midnight Bat-Tortoise" Edition

Converte HTML estruturado com CSS inline (cores, fontes, tamanhos, alinhamentos,
tabelas, listas, imagens e margens) para DOCX nativo com fidelidade milimétrica.
"""

import os
import re
import io
import base64
import logging
import tempfile
from typing import Optional
from urllib.parse import unquote

from fastapi import FastAPI, HTTPException, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from bs4 import BeautifulSoup, Tag, NavigableString

import docx
from docx.shared import Mm, Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

import pypandoc

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)
logger = logging.getLogger("webdoc-docx")

app = FastAPI(
    title="WebDoc — DOCX Export Service (High Fidelity)",
    description="Microsserviço FastAPI para conversão de HTML para DOCX com preservação integral de formatação visual.",
    version="2.0.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

MEDIA_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "../../media-library"))


class ExportRequest(BaseModel):
    html: str = Field(..., description="Conteúdo HTML completo ou parcial do documento")
    docName: Optional[str] = Field(None, description="Nome base do arquivo gerado")
    doc_name: Optional[str] = Field(None, description="Nome base do arquivo gerado (compatibilidade)")
    pageWidth: Optional[float] = Field(None, description="Largura da página em mm")
    pageHeight: Optional[float] = Field(None, description="Altura da página em mm")
    page_width_mm: Optional[float] = Field(210.0, description="Largura da página em mm (padrão: A4 = 210mm)")
    page_height_mm: Optional[float] = Field(297.0, description="Altura da página em mm (padrão: A4 = 297mm)")
    marginTop: Optional[float] = Field(None, description="Margem superior em mm")
    marginBottom: Optional[float] = Field(None, description="Margem inferior em mm")
    marginLeft: Optional[float] = Field(None, description="Margem esquerda em mm")
    marginRight: Optional[float] = Field(None, description="Margem direita em mm")
    margin_top_mm: Optional[float] = Field(None, description="Margem superior em mm")
    margin_bottom_mm: Optional[float] = Field(None, description="Margem inferior em mm")
    margin_left_mm: Optional[float] = Field(None, description="Margem esquerda em mm")
    margin_right_mm: Optional[float] = Field(None, description="Margem direita em mm")
    landscape: Optional[bool] = Field(False, description="Orientação paisagem")
    format_name: Optional[str] = Field("A4", description="Nome do formato (informativo)")
    formatName: Optional[str] = Field(None, description="Nome do formato (informativo)")
    pageMarginsMap: Optional[dict] = Field(None, description="Mapa de margens por folha")


def parse_color(color_str: str) -> tuple[int, int, int] | None:
    if not color_str:
        return None
    c = color_str.strip().lower()
    if c.startswith("#"):
        hex_val = c[1:]
        if len(hex_val) == 3:
            hex_val = "".join([ch * 2 for ch in hex_val])
        if len(hex_val) == 6:
            try:
                return (int(hex_val[0:2], 16), int(hex_val[2:4], 16), int(hex_val[4:6], 16))
            except ValueError:
                return None
    rgb_match = re.match(r"rgba?\s*\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)", c)
    if rgb_match:
        return (int(rgb_match.group(1)), int(rgb_match.group(2)), int(rgb_match.group(3)))
    hsl_match = re.match(r"hsla?\s*\(\s*(\d+)\s*,\s*(\d+)%?\s*,\s*(\d+)%", c)
    if hsl_match:
        import colorsys
        h = float(hsl_match.group(1)) / 360.0
        s = float(hsl_match.group(2)) / 100.0
        l = float(hsl_match.group(3)) / 100.0
        r, g, b = colorsys.hls_to_rgb(h, l, s)
        return (int(r * 255), int(g * 255), int(b * 255))
    NAMED_COLORS = {
        "black": (0, 0, 0), "white": (255, 255, 255), "red": (220, 53, 69), "green": (40, 167, 69),
        "blue": (0, 123, 255), "gray": (128, 128, 128), "grey": (128, 128, 128), "orange": (217, 93, 57),
        "purple": (107, 62, 117), "dark": (17, 17, 17), "light": (243, 233, 210)
    }
    return NAMED_COLORS.get(c)


def parse_size_pt(size_str: str) -> float | None:
    if not size_str:
        return None
    s = size_str.strip().lower()
    m = re.match(r"([\d.]+)\s*(pt|px|em|rem)?", s)
    if not m:
        return None
    val = float(m.group(1))
    unit = m.group(2) or "px"
    if unit == "pt":
        return val
    if unit == "px":
        return val * 0.75
    if unit in ("em", "rem"):
        return val * 11.0
    return val


def parse_inline_styles(style_str: str) -> dict[str, str]:
    styles = {}
    if not style_str:
        return styles
    for part in style_str.split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            styles[k.strip().lower()] = v.strip()
    return styles


def clean_font_family(font_str: str) -> str | None:
    if not font_str or font_str.strip() == "default":
        return None
    primary = font_str.split(",")[0].strip().replace("'", "").replace('"', "")
    return primary if primary else None


def set_cell_background(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_run_background(run, hex_color: str):
    rPr = run._r.get_or_add_rPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    rPr.append(shd)


def set_run_font_name(run, font_name: str):
    if not font_name:
        return
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:cs"), font_name)


def add_hyperlink(paragraph, url: str, text: str, style_dict: dict):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{r_id}"/>'
    )
    new_run = parse_xml(f'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:t>{text}</w:t></w:r>')
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

    rPr = new_run.get_or_add_rPr()
    color_hex = "D95D39"
    if style_dict.get("color"):
        rgb = parse_color(style_dict["color"])
        if rgb:
            color_hex = f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
    rColor = parse_xml(f'<w:color xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="{color_hex}"/>')
    rPr.append(rColor)
    u = parse_xml('<w:u xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="single"/>')
    rPr.append(u)


def resolve_image_data(src: str, media_root: str) -> bytes | None:
    if not src:
        return None
    if src.startswith("data:"):
        try:
            _, encoded = src.split(",", 1)
            return base64.b64decode(encoded)
        except Exception:
            return None
    if "/media/" in src:
        try:
            sub = src.split("/media/")[-1]
            decoded = unquote(sub.split("?")[0])
            local_path = os.path.join(media_root, decoded)
            if os.path.exists(local_path) and os.path.isfile(local_path):
                with open(local_path, "rb") as f:
                    return f.read()
        except Exception:
            pass
    if os.path.exists(src) and os.path.isfile(src):
        with open(src, "rb") as f:
            return f.read()
    return None


class RichDocxBuilder:
    def __init__(
        self,
        page_width_mm=210,
        page_height_mm=297,
        landscape=False,
        margin_top_mm=25,
        margin_bottom_mm=25,
        margin_left_mm=20,
        margin_right_mm=20,
        page_margins_map=None,
        media_root=None,
    ):
        self.doc = docx.Document()
        self.media_root = media_root or MEDIA_ROOT
        self.page_margins_map = page_margins_map or {}
        self.current_page = 1

        self.w = max(10.0, float(page_width_mm))
        self.h = max(10.0, float(page_height_mm))
        self.landscape = bool(landscape)

        if self.landscape:
            self.w, self.h = max(self.w, self.h), min(self.w, self.h)
            self.orientation = WD_ORIENT.LANDSCAPE
        else:
            self.w, self.h = min(self.w, self.h), max(self.w, self.h)
            self.orientation = WD_ORIENT.PORTRAIT

        self._apply_section_properties(
            self.doc.sections[0],
            margin_top_mm,
            margin_bottom_mm,
            margin_left_mm,
            margin_right_mm
        )

    def _apply_section_properties(self, section, top, btm, lft, rgt):
        section.orientation = self.orientation
        section.page_width = Mm(self.w)
        section.page_height = Mm(self.h)
        section.top_margin = Mm(max(0.0, float(top)))
        section.bottom_margin = Mm(max(0.0, float(btm)))
        section.left_margin = Mm(max(0.0, float(lft)))
        section.right_margin = Mm(max(0.0, float(rgt)))
        section.header_margin = Mm(10.0)
        section.footer_margin = Mm(10.0)

    def convert_html(self, html_content: str) -> docx.Document:
        soup = BeautifulSoup(html_content, "html.parser")

        # Limpa elementos de UI do editor
        for el in soup.find_all(
            class_=re.compile(r"multi-page-break-label|page-guide-box|page-boundary-marker|img-resizer-overlay|link-preview-balloon")
        ):
            el.decompose()

        body = soup.body if soup.body else soup
        for child in body.children:
            if isinstance(child, NavigableString):
                text = str(child).strip()
                if text:
                    p = self.doc.add_paragraph()
                    self._apply_run_styles(p.add_run(text), {})
            elif isinstance(child, Tag):
                self._process_block_element(child)

        return self.doc

    def _process_block_element(self, tag: Tag):
        tag_name = tag.name.lower()

        # Quebras de página (com suporte a nova seção caso a próxima folha tenha margem personalizada)
        if "page-break" in tag.get("class", []) or "multi-page-break" in tag.get("class", []):
            self.current_page += 1
            next_m = self.page_margins_map.get(self.current_page) or self.page_margins_map.get(str(self.current_page))
            if next_m and isinstance(next_m, dict) and "top" in next_m:
                new_sec = self.doc.add_section()
                self._apply_section_properties(
                    new_sec,
                    next_m.get("top", 25),
                    next_m.get("bottom", 25),
                    next_m.get("left", 20),
                    next_m.get("right", 20)
                )
            else:
                self.doc.add_page_break()
            return

        # Títulos (H1-H6)
        if tag_name in ["h1", "h2", "h3", "h4", "h5", "h6"]:
            level = int(tag_name[1])
            p = self.doc.add_paragraph()
            self._apply_paragraph_styles(p, tag, default_space_before=14 - level * 2, default_space_after=6)

            tag_styles = parse_inline_styles(tag.get("style", ""))
            default_sizes = {1: 24, 2: 18, 3: 14, 4: 12, 5: 11, 6: 10}
            default_color = "#D95D39" if level == 3 else "#111111"

            h_styles = {
                "font-family": tag_styles.get("font-family") or "Bangers",
                "font-size": tag_styles.get("font-size") or f"{default_sizes.get(level, 12)}pt",
                "bold": True,
                "color": tag_styles.get("color") or default_color,
            }

            self._process_inline_children(tag, p, inherited_styles=h_styles)
            return

        # Parágrafos
        if tag_name == "p":
            p = self.doc.add_paragraph()
            self._apply_paragraph_styles(p, tag, default_space_after=6)
            
            tag_styles = parse_inline_styles(tag.get("style", ""))
            p_inherited = {}
            if tag_styles.get("font-family"):
                p_inherited["font-family"] = tag_styles["font-family"]
            if tag_styles.get("font-size"):
                p_inherited["font-size"] = tag_styles["font-size"]
            if tag_styles.get("color"):
                p_inherited["color"] = tag_styles["color"]

            self._process_inline_children(tag, p, inherited_styles=p_inherited)
            return

        # Citações / Blockquote
        if tag_name == "blockquote":
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            self._apply_paragraph_styles(p, tag)
            self._process_inline_children(tag, p, inherited_styles={"italic": True, "color": "#555555"})
            return

        # Bloco de código
        if tag_name in ["pre", "code"]:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Inches(0.2)
            self._process_inline_children(
                tag, p, inherited_styles={"font-family": "JetBrains Mono", "font-size": "9.5pt", "color": "#111111"}
            )
            return

        # Listas (UL / OL)
        if tag_name in ["ul", "ol"]:
            self._process_list(tag, is_ordered=(tag_name == "ol"))
            return

        # Tabelas
        if tag_name == "table":
            self._process_table(tag)
            return

        # Imagens
        if tag_name == "figure" or tag_name == "img":
            self._process_image(tag)
            return

        # Divisor horizontal
        if tag_name == "hr":
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            r = p.add_run("―" * 40)
            r.font.color.rgb = RGBColor(17, 17, 17)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return

        # Containers genéricos (div, section, etc.)
        for child in tag.children:
            if isinstance(child, Tag):
                self._process_block_element(child)
            elif isinstance(child, NavigableString) and str(child).strip():
                p = self.doc.add_paragraph()
                self._apply_run_styles(p.add_run(str(child).strip()), {})

    def _apply_paragraph_styles(self, p, tag: Tag, default_space_before=0, default_space_after=4):
        styles = parse_inline_styles(tag.get("style", ""))
        classes = tag.get("class", [])

        align = styles.get("text-align") or ""
        if "text-align-center" in classes or align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif "text-align-right" in classes or align == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif "text-align-justify" in classes or align == "justify":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif "text-align-left" in classes or align == "left":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        p.paragraph_format.space_before = Pt(default_space_before)
        p.paragraph_format.space_after = Pt(default_space_after)
        p.paragraph_format.line_spacing = 1.25

    def _process_list(self, list_tag: Tag, is_ordered=False, level=0):
        for li in list_tag.find_all("li", recursive=False):
            p = self.doc.add_paragraph(style="List Number" if is_ordered else "List Bullet")
            p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
            p.paragraph_format.space_after = Pt(3)
            self._process_inline_children(li, p)

            for sub_list in li.find_all(["ul", "ol"], recursive=False):
                self._process_list(sub_list, is_ordered=(sub_list.name == "ol"), level=level + 1)

    def _process_table(self, table_tag: Tag):
        rows = table_tag.find_all("tr")
        if not rows:
            return

        col_count = max(len(r.find_all(["th", "td"])) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=col_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        for row_idx, r in enumerate(rows):
            cells = r.find_all(["th", "td"])
            for col_idx, cell_tag in enumerate(cells):
                if col_idx >= col_count:
                    break
                doc_cell = table.cell(row_idx, col_idx)
                doc_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                is_th = cell_tag.name == "th"
                styles = parse_inline_styles(cell_tag.get("style", ""))

                bg_color = styles.get("background-color") or styles.get("background")
                if bg_color:
                    rgb = parse_color(bg_color)
                    if rgb:
                        set_cell_background(doc_cell, f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}")
                elif is_th:
                    set_cell_background(doc_cell, "111111")

                p = doc_cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)

                inherited = {}
                if is_th:
                    inherited = {"bold": True, "color": "#F3E9D2"}
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                self._process_inline_children(cell_tag, p, inherited_styles=inherited)

        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def _process_image(self, tag: Tag):
        img_tag = tag if tag.name == "img" else tag.find("img")
        if not img_tag:
            return

        src = img_tag.get("src", "")
        img_bytes = resolve_image_data(src, self.media_root)
        if not img_bytes:
            return

        p = self.doc.add_paragraph()
        classes = tag.get("class", []) + img_tag.get("class", [])

        if "image-style-align-left" in classes or "image-style-side" in classes:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif "image-style-align-right" in classes:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)

        styles = parse_inline_styles(img_tag.get("style", ""))
        width_str = styles.get("width") or img_tag.get("width")
        width_val = None
        if width_str:
            m = re.match(r"([\d.]+)\s*(px|mm|cm|%)?", str(width_str))
            if m:
                v = float(m.group(1))
                u = m.group(2) or "px"
                if u == "mm":
                    width_val = Mm(v)
                elif u == "cm":
                    width_val = Mm(v * 10)
                elif u == "px":
                    width_val = Inches(v / 96.0)
                elif u == "%":
                    width_val = Inches(5.5 * (v / 100.0))

        if not width_val:
            width_val = Inches(5.0)

        try:
            stream = io.BytesIO(img_bytes)
            run = p.add_run()
            run.add_picture(stream, width=width_val)
        except Exception as e:
            logger.warning(f"Erro ao embutir imagem no DOCX: {e}")

    def _process_inline_children(self, parent_tag: Tag, paragraph, inherited_styles=None):
        inherited = dict(inherited_styles or {})

        for child in parent_tag.children:
            if isinstance(child, NavigableString):
                text = str(child)
                if text:
                    run = paragraph.add_run(text)
                    self._apply_run_styles(run, inherited)
            elif isinstance(child, Tag):
                child_name = child.name.lower()
                child_styles = dict(inherited)

                tag_styles = parse_inline_styles(child.get("style", ""))
                child_classes = child.get("class", [])

                if tag_styles.get("color"):
                    child_styles["color"] = tag_styles["color"]
                if tag_styles.get("background-color") or tag_styles.get("background"):
                    child_styles["background-color"] = tag_styles.get("background-color") or tag_styles.get("background")
                if tag_styles.get("font-family"):
                    child_styles["font-family"] = tag_styles["font-family"]
                if tag_styles.get("font-size"):
                    child_styles["font-size"] = tag_styles["font-size"]

                if "text-tiny" in child_classes:
                    child_styles["font-size"] = "8.5pt"
                elif "text-small" in child_classes:
                    child_styles["font-size"] = "10pt"
                elif "text-big" in child_classes:
                    child_styles["font-size"] = "14pt"
                elif "text-huge" in child_classes:
                    child_styles["font-size"] = "18pt"

                if child_name in ["strong", "b"] or tag_styles.get("font-weight") in ["bold", "700", "800", "900"]:
                    child_styles["bold"] = True
                if child_name in ["em", "i"] or tag_styles.get("font-style") == "italic":
                    child_styles["italic"] = True
                if child_name == "u" or "underline" in tag_styles.get("text-decoration", ""):
                    child_styles["underline"] = True
                if child_name in ["s", "strike", "del"] or "line-through" in tag_styles.get("text-decoration", ""):
                    child_styles["strike"] = True
                if child_name == "sub":
                    child_styles["subscript"] = True
                if child_name == "sup":
                    child_styles["superscript"] = True
                if child_name == "mark":
                    child_styles["background-color"] = "#FFEB3B"
                if child_name == "code":
                    child_styles["font-family"] = "JetBrains Mono"
                    child_styles["color"] = "#D95D39"
                    child_styles["background-color"] = "#F4EBE1"

                if child_name == "a" and child.get("href"):
                    add_hyperlink(paragraph, child.get("href"), child.get_text(), child_styles)
                    continue

                if child_name == "br":
                    paragraph.add_run("\n")
                    continue

                self._process_inline_children(child, paragraph, child_styles)

    def _apply_run_styles(self, run, styles: dict):
        if styles.get("bold"):
            run.bold = True
        if styles.get("italic"):
            run.italic = True
        if styles.get("underline"):
            run.underline = True
        if styles.get("strike"):
            run.font.strike = True
        if styles.get("subscript"):
            run.font.subscript = True
        if styles.get("superscript"):
            run.font.superscript = True

        font_fam = clean_font_family(styles.get("font-family"))
        if font_fam:
            set_run_font_name(run, font_fam)
        else:
            set_run_font_name(run, "Merriweather")

        size_pt = parse_size_pt(styles.get("font-size"))
        if size_pt:
            run.font.size = Pt(size_pt)
        else:
            run.font.size = Pt(11)

        color_str = styles.get("color")
        if color_str:
            rgb = parse_color(color_str)
            if rgb:
                run.font.color.rgb = RGBColor(rgb[0], rgb[1], rgb[2])
        else:
            run.font.color.rgb = RGBColor(17, 17, 17)

        bg_str = styles.get("background-color")
        if bg_str:
            rgb = parse_color(bg_str)
            if rgb:
                set_run_background(run, f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}")


@app.post(
    "/api/export/docx",
    summary="Exportar documento para DOCX com alta fidelidade visual",
    response_description="Arquivo .docx pronto para download",
)
async def export_docx(req: ExportRequest):
    if not req.html or not req.html.strip():
        raise HTTPException(status_code=400, detail="O campo 'html' não pode estar vazio.")

    raw_w = req.pageWidth if req.pageWidth is not None else (req.page_width_mm or 210.0)
    raw_h = req.pageHeight if req.pageHeight is not None else (req.page_height_mm or 297.0)

    # Extrai margens com suporte total a pageMarginsMap editado pelo usuário
    page_map = req.pageMarginsMap or {}
    p1 = page_map.get(1) or page_map.get("1") or {}

    top_cand = p1.get("top") if p1.get("top") is not None else (req.marginTop if req.marginTop is not None else req.margin_top_mm)
    btm_cand = p1.get("bottom") if p1.get("bottom") is not None else (req.marginBottom if req.marginBottom is not None else req.margin_bottom_mm)
    lft_cand = p1.get("left") if p1.get("left") is not None else (req.marginLeft if req.marginLeft is not None else req.margin_left_mm)
    rgt_cand = p1.get("right") if p1.get("right") is not None else (req.marginRight if req.marginRight is not None else req.margin_right_mm)

    top_val = float(top_cand) if top_cand is not None else 25.0
    btm_val = float(btm_cand) if btm_cand is not None else 25.0
    lft_val = float(lft_cand) if lft_cand is not None else 20.0
    rgt_val = float(rgt_cand) if rgt_cand is not None else 20.0

    doc_title = req.docName or req.doc_name or "documento"
    format_name = req.formatName or req.format_name or "A4"

    logger.info(
        f"Exportando DOCX | Documento: {doc_title} | Formato: {format_name} | "
        f"{raw_w}x{raw_h}mm | Margens: {top_val}/{btm_val}/{lft_val}/{rgt_val}mm | Paisagem: {req.landscape}"
    )

    try:
        builder = RichDocxBuilder(
            page_width_mm=raw_w,
            page_height_mm=raw_h,
            landscape=bool(req.landscape),
            margin_top_mm=top_val,
            margin_bottom_mm=btm_val,
            margin_left_mm=lft_val,
            margin_right_mm=rgt_val,
            page_margins_map=page_map,
            media_root=MEDIA_ROOT,
        )
        doc = builder.convert_html(req.html)

        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        safe_name = str(doc_title).replace(" ", "_")
        safe_name = "".join(c for c in safe_name if c.isalnum() or c in "-_")[:50] or "documento"
        safe_fmt = str(format_name).replace(" ", "_")
        safe_fmt = "".join(c for c in safe_fmt if c.isalnum() or c in "-_")[:30] or "A4"
        filename = f"{safe_name}_{safe_fmt}.docx"

        logger.info(f"DOCX de alta fidelidade gerado com sucesso: {len(docx_bytes)} bytes")

        return Response(
            content=docx_bytes,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Access-Control-Expose-Headers": "Content-Disposition, Content-Type",
            },
        )

    except Exception as e:
        logger.error(f"Erro no RichDocxBuilder: {e}. Tentando fallback Pandoc...", exc_info=True)
        try:
            with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
                tmp_path = tmp.name

            pypandoc.convert_text(
                source=req.html,
                to="docx",
                format="html",
                outputfile=tmp_path,
            )
            with open(tmp_path, "rb") as f:
                docx_bytes = f.read()

            if os.path.exists(tmp_path):
                os.unlink(tmp_path)

            return Response(
                content=docx_bytes,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={
                    "Content-Disposition": f'attachment; filename="{doc_title}.docx"',
                    "Access-Control-Expose-Headers": "Content-Disposition, Content-Type",
                },
            )
        except Exception as pandoc_err:
            logger.error(f"Falha no fallback Pandoc: {pandoc_err}")
            raise HTTPException(status_code=500, detail=f"Erro na conversão: {str(e)}")


def _get_pandoc_status() -> str:
    try:
        return pypandoc.get_pandoc_version()
    except Exception:
        return "unavailable"


@app.get("/health", summary="Status do serviço de exportação DOCX")
async def health():
    return {
        "status": "ok",
        "service": "WebDoc Rich DOCX Export",
        "engine": "RichDocxBuilder 2.0 (python-docx + bs4)",
        "pandoc_version": _get_pandoc_status(),
    }


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000, reload=True)
